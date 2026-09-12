"""JSON tables and CSV/XLSX interchange. Row indices are zero based."""
import csv
import base64
import io
import math
import re
from autoflow.datatypes import SCALARS, canonical, convert
from datetime import date, datetime, time


class Dataset(dict):
    """Shared table object. Mutations preserve identity across graph connections."""


def get_dataset(project, name):
    tables = project.setdefault('datasets', {})
    if name not in tables: raise ValueError(f'Unknown dataset: {name}')
    if not isinstance(tables[name], Dataset): tables[name] = Dataset(validate(tables[name]))
    return tables[name]


def infer_type(values):
    kinds = set()
    for value in values:
        if value is None or value == '': continue
        if type(value) is bool: kind = 'bool'
        elif type(value) is int: kind = 'integer'
        elif type(value) is float: kind = 'float'
        elif type(value) is bytes: kind = 'bytes'
        elif isinstance(value, str) and value.lower() in ('true', 'false'): kind = 'bool'
        elif isinstance(value, str) and re.fullmatch(r'-?(0|[1-9]\d*)(\.\d+)?([eE][+-]?\d+)?', value):
            kind = ('float' if any(c in value.lower() for c in '.e') else 'integer') if math.isfinite(float(value)) else 'string'
        else: kind = 'string'
        kinds.add(kind)
    if kinds and kinds <= {'integer', 'float'}: return 'float' if 'float' in kinds else 'integer'
    return next(iter(kinds)) if len(kinds) == 1 else 'string'


def validate(table):
    if not isinstance(table, dict): raise ValueError('Dataset must be an object')
    columns, rows = table.get('columns'), table.get('rows')
    if not isinstance(columns, list) or not all(isinstance(c, str) and c for c in columns) or len(set(columns)) != len(columns):
        raise ValueError('Columns must be unique nonempty names')
    if not isinstance(rows, list): raise ValueError('Rows must be an array')
    for row in rows:
        if not isinstance(row, dict) or set(row) != set(columns): raise ValueError('Each row must contain exactly the dataset columns')
        for value in row.values():
            if value is not None and type(value) not in (str, int, float, bool, bytes): raise ValueError('Dataset cells must be bool, integer, float, string or bytes')
            if type(value) is float and not math.isfinite(value): raise ValueError('Numbers must be finite')
    types = table.get('types', {})
    if not isinstance(types, dict) or set(types) - set(columns): raise ValueError('Invalid column types')
    normalized = []
    resolved = {}
    for c in columns:
        kind = canonical(types[c]) if c in types else infer_type([row[c] for row in rows])
        if kind not in SCALARS: raise ValueError(f'Unsupported dataset type: {kind}')
        resolved[c] = kind
    for row in rows:
        normalized.append({c: convert(row[c], resolved[c]) if row[c] is not None and row[c] != '' else
                           (False if resolved[c] == 'bool' else 0 if resolved[c] == 'integer' else 0.0 if resolved[c] == 'float' else b'' if resolved[c] == 'bytes' else '') for c in columns})
    table['types'], table['rows'] = resolved, normalized
    return table


def validate_all(datasets):
    if not isinstance(datasets, dict): raise ValueError('Datasets must be an object')
    for name, table in datasets.items():
        if not isinstance(name, str) or not name.strip() or name in ('__proto__', 'constructor', 'prototype'): raise ValueError('Invalid dataset name')
        validate(table)


def read(raw, extension, sheet=None):
    sheets = []
    if extension.lower() == '.csv':
        text = raw.decode('utf-8-sig')
        try: dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        values = list(csv.reader(io.StringIO(text), dialect))
    elif extension.lower() == '.xlsx':
        from openpyxl import load_workbook
        from zipfile import ZipFile
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100_000_000: raise ValueError('Excel file exceeds 100 MB expanded')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            sheets = workbook.sheetnames
            values = list(workbook[sheet or sheets[0]].iter_rows(values_only=True))
        finally: workbook.close()
    else: raise ValueError('Choose CSV or Excel .xlsx')
    if not values: raise ValueError('File has no header row')
    columns = [str(c) if c is not None else '' for c in values[0]]
    rows = []
    for cells in values[1:]:
        if len(cells) > len(columns): raise ValueError('Row has more cells than headers')
        cells = list(cells) + [None] * (len(columns) - len(cells))
        cells = [v[len('autoflow:string::'):] if isinstance(v, str) and v.startswith('autoflow:string::') else base64.b64decode(v[len('base64:bytes::'):], validate=True) if isinstance(v, str) and v.startswith('base64:bytes::') else v for v in cells]
        rows.append(dict(zip(columns, [v.isoformat() if isinstance(v, (date, datetime, time)) else v for v in cells])))
    return validate(dict(columns=columns, rows=rows)), sheets


def write(table, extension):
    validate(table)
    cell = lambda v: 'base64:bytes::' + base64.b64encode(v).decode('ascii') if isinstance(v, bytes) else 'autoflow:string::' + v if isinstance(v, str) and v.startswith(('base64:bytes::', 'autoflow:string::')) else v
    if extension == '.csv':
        stream = io.StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(table['columns'])
        writer.writerows([cell(r[c]) for c in table['columns']] for r in table['rows'])
        return stream.getvalue().encode('utf-8-sig')
    if extension != '.xlsx': raise ValueError('Choose CSV or Excel .xlsx')
    from openpyxl import Workbook
    workbook = Workbook()
    ws = workbook.active
    ws.title = 'Data'
    ws.append(table['columns'])
    for row in table['rows']: ws.append([cell(row[c]) for c in table['columns']])
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def row_at(table, index):
    if type(index) is not int or not 0 <= index < len(table['rows']): raise ValueError('Row index out of range (indices start at 0)')
    return table['rows'][index]


def column(table, name):
    if name not in table['columns']: raise ValueError(f'Unknown column: {name}')
    return name


def sort(table, name, descending=False):
    column(table, name)
    def key(row):
        v = row[name]
        return (0, '') if v is None else (1, v) if type(v) in (int, float) else (2, str(v).casefold())
    table['rows'].sort(key=key, reverse=descending)
